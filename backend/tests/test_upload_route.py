"""TDD: POST /api/upload (multipart PDF -> Document row, queued for the worker)."""

import io
from datetime import datetime, timezone
from decimal import Decimal

import pytest
from fastapi import HTTPException

from config import settings
from contracts import TopicProfile
from db.models import Document, Session as SessionModel, User, UsageCounter
from routes.upload import _read_bounded
from services import cost_meter


SESSION_ID = "sess_up"
USER_ID = "u_up"


@pytest.fixture
def seeded(db_session):
    db_session.add(User(id=USER_ID))
    db_session.flush()
    db_session.add(
        SessionModel(
            id=SESSION_ID,
            user_id=USER_ID,
            topic="sql",
            topic_profile_json=TopicProfile().model_dump_json(),
        )
    )
    db_session.commit()


@pytest.fixture(autouse=True)
def ingestion_never_inline(monkeypatch):
    """F-04: ingestion must be enqueued for the worker, never run in-process."""

    def _boom(doc_id):
        raise AssertionError("ingestion must not run in the request process")

    monkeypatch.setattr("services.ingestion_service.run", _boom)


@pytest.fixture(autouse=True)
def stub_filesystem(monkeypatch, tmp_path):
    monkeypatch.setattr("routes.upload.settings.uploads_path", str(tmp_path))


def test_upload_returns_202_and_creates_pending_document(client, seeded, db_session):
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 202, r.text
    body = r.json()
    assert body["status"] == "pending"
    assert body["filename"] == "notes.pdf"
    assert body["session_id"] == SESSION_ID

    doc = db_session.get(Document, body["document_id"])
    assert doc is not None
    assert doc.status == "pending"


def test_disallowed_extension_400(client, seeded):
    files = {"file": ("paper.docx", io.BytesIO(b"PK\x03\x04"), "application/octet-stream")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 400


def test_extensionless_filename_400(client, seeded):
    files = {"file": ("README", io.BytesIO(b"hello"), "text/plain")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 400


@pytest.mark.parametrize(
    "name,ctype,content",
    [
        (
            "slides.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            b"PK\x03\x04data-bytes",
        ),
        ("notes.txt", "text/plain", b"data-bytes"),
        ("notes.md", "text/markdown", b"data-bytes"),
        ("notes.markdown", "application/octet-stream", b"data-bytes"),
    ],
)
def test_allowed_non_pdf_types_202(client, seeded, name, ctype, content):
    files = {"file": (name, io.BytesIO(content), ctype)}
    r = client.post(
        "/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files
    )
    assert r.status_code == 202, r.text
    assert r.json()["status"] == "pending"


# ---------------------------------------------------------------------------
# I-03: X-Cost-Warning header on soft-cap breach (mirrors test_cost_cap.py's
# end_session coverage for the other non-SSE response that needs the header).
# ---------------------------------------------------------------------------


def test_upload_sets_cost_warning_header_when_soft_breached(client, seeded, db_session):
    cost_meter.record_cost(
        db_session, USER_ID, Decimal(str(settings.llm_soft_cap_usd)) + Decimal("0.01")
    )
    db_session.commit()

    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 202, r.text
    assert r.headers["x-cost-warning"].startswith("level=")


def test_upload_no_header_under_soft_cap(client, seeded):
    files = {"file": ("notes2.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 202, r.text
    assert "x-cost-warning" not in r.headers


def test_missing_session_id_field_400(client, seeded):
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", files=files)
    assert r.status_code in (400, 422)


def test_unknown_session_id_404(client, seeded):
    # Phase 7: ownership check folds unknown-session and wrong-owner into 404
    # to avoid an existence oracle (matches the sessions/profile pattern).
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": "does_not_exist"},
        files=files,
    )
    assert r.status_code == 404


def test_upload_into_ended_session_409(client, seeded, db_session):
    """C-07: an ended session is read-only. The 409 must land before the cost
    gate and the rate limiter, so it burns no daily slot and writes no row."""
    sess = db_session.get(SessionModel, SESSION_ID)
    sess.ended_at = datetime.now(timezone.utc)
    db_session.commit()

    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post(
        "/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files
    )
    assert r.status_code == 409, r.text
    assert r.json()["detail"]["code"] == "session_ended"
    assert db_session.query(Document).count() == 0
    assert db_session.query(UsageCounter).filter_by(user_id=USER_ID).count() == 0


def test_missing_auth_header_401(client, seeded):
    # Phase 7: user_id no longer carried in form. Missing Authorization
    # header => 401. Override returns "test-user" default => /api/upload
    # would reach ownership check and 404 because seeded session belongs to
    # a different user. Either response is acceptable as long as auth gates
    # the request before the form is even validated.
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"session_id": SESSION_ID}, files=files)
    assert r.status_code in (401, 404)


def test_upload_returns_429_when_cap_reached(client, seeded, monkeypatch):
    monkeypatch.setattr(
        "routes.upload.rate_limit.check_and_increment",
        lambda db, uid: (False, 9999),
    )
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    assert r.status_code == 429
    body = r.json()["detail"]
    assert body["code"] == "daily_cap_reached"
    assert body["used"] == 9999
    assert "resets_at" in body


def test_get_upload_status_returns_current_doc_state(client, seeded, db_session):
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    doc_id = r.json()["document_id"]

    r2 = client.get(f"/api/upload/{doc_id}?user_id={USER_ID}")
    assert r2.status_code == 200, r2.text
    body = r2.json()
    assert body["id"] == doc_id
    assert body["status"] == "pending"
    assert body["error"] is None

    doc = db_session.get(Document, doc_id)
    doc.status = "failed"
    doc.error = "embedding service unreachable"
    db_session.commit()

    r3 = client.get(f"/api/upload/{doc_id}?user_id={USER_ID}")
    assert r3.status_code == 200
    body3 = r3.json()
    assert body3["status"] == "failed"
    assert body3["error"] == "embedding service unreachable"


def test_get_upload_status_404_for_missing(client):
    r = client.get(f"/api/upload/99999?user_id={USER_ID}")
    assert r.status_code == 404


def test_get_upload_status_404_for_wrong_user(client, seeded, db_session):
    db_session.add(User(id="other"))
    db_session.commit()
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    doc_id = r.json()["document_id"]
    r2 = client.get(f"/api/upload/{doc_id}?user_id=other")
    assert r2.status_code == 404


def test_upload_rejects_oversize_via_content_length(client, seeded):
    big = b"%PDF-" + b"x" * (26 * 1024 * 1024)
    files = {"file": ("big.pdf", io.BytesIO(big), "application/pdf")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    assert r.status_code == 413
    assert r.json()["detail"]["code"] == "FILE_TOO_LARGE"


def test_upload_sanitizes_traversal_filename(client, seeded, db_session):
    files = {"file": ("../../etc/passwd.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    assert r.status_code == 202, r.text
    body = r.json()
    assert "/" not in body["filename"]
    assert "\\" not in body["filename"]
    assert ".." not in body["filename"]
    assert body["filename"].endswith("passwd.pdf")


def test_upload_leaves_document_pending_for_worker_to_claim(client, seeded, db_session):
    """F-04: upload only enqueues (status="pending"); the worker process
    claims and runs ingestion out-of-process (see services/worker.py)."""
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 202
    doc = db_session.get(Document, r.json()["document_id"])
    assert doc.status == "pending"


# ---------------------------------------------------------------------------
# C-08: identical re-upload into the same session is idempotent.
# ---------------------------------------------------------------------------

PDF_BYTES = b"%PDF-dedupe-fixture"


def _post_pdf(client, session_id=SESSION_ID, content=PDF_BYTES, name="notes.pdf"):
    return client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": session_id},
        files={"file": (name, io.BytesIO(content), "application/pdf")},
    )


def test_identical_reupload_returns_same_document_id(client, seeded, db_session, monkeypatch):
    puts = []

    class RecordingStore:
        def put(self, key, data):
            puts.append(key)

    monkeypatch.setattr("routes.upload.object_store.get_store", lambda: RecordingStore())

    first = _post_pdf(client)
    assert first.status_code == 202, first.text
    second = _post_pdf(client)
    assert second.status_code == 202, second.text

    assert second.json()["document_id"] == first.json()["document_id"]
    assert db_session.query(Document).count() == 1
    assert len(puts) == 1


def test_identical_bytes_in_another_session_create_a_second_row(client, seeded, db_session):
    other = "sess_up_2"
    db_session.add(
        SessionModel(
            id=other,
            user_id=USER_ID,
            topic="sql-2",
            topic_profile_json=TopicProfile().model_dump_json(),
        )
    )
    db_session.commit()

    first = _post_pdf(client)
    second = _post_pdf(client, session_id=other)
    assert first.status_code == 202 and second.status_code == 202, second.text
    assert first.json()["document_id"] != second.json()["document_id"]
    assert db_session.query(Document).count() == 2


def test_failed_row_does_not_block_a_retry(client, seeded, db_session):
    import hashlib

    db_session.add(
        Document(
            session_id=SESSION_ID,
            filename="notes.pdf",
            status="failed",
            error="storage write failed",
            content_sha256=hashlib.sha256(PDF_BYTES).hexdigest(),
        )
    )
    db_session.commit()

    r = _post_pdf(client)
    assert r.status_code == 202, r.text
    doc = db_session.get(Document, r.json()["document_id"])
    assert doc.status == "pending"
    assert db_session.query(Document).count() == 2


def test_upload_populates_content_sha256(client, seeded, db_session):
    import hashlib

    r = _post_pdf(client)
    assert r.status_code == 202, r.text
    doc = db_session.get(Document, r.json()["document_id"])
    assert doc.content_sha256 == hashlib.sha256(PDF_BYTES).hexdigest()


def test_identical_reupload_does_not_burn_a_rate_limit_slot(client, seeded, db_session):
    assert _post_pdf(client).status_code == 202
    counter = db_session.query(UsageCounter).filter_by(user_id=USER_ID).one()
    assert counter.count == 1

    assert _post_pdf(client).status_code == 202
    db_session.refresh(counter)
    assert counter.count == 1


def test_concurrent_identical_uploads_resolve_to_the_existing_row(
    client, seeded, db_session, monkeypatch
):
    """C-08 race: another request commits the same (session_id, sha) between
    our lookup and our flush. The partial unique index raises IntegrityError;
    we roll back, re-select, and return the winner's payload."""
    import hashlib

    from routes import upload as upload_module

    real_lookup = upload_module._find_existing_document
    calls = {"n": 0}

    def racing_lookup(db, session_id, sha):
        calls["n"] += 1
        if calls["n"] == 1:
            db.add(
                Document(
                    session_id=session_id,
                    filename="winner.pdf",
                    status="pending",
                    content_sha256=sha,
                )
            )
            db.commit()
            return None
        return real_lookup(db, session_id, sha)

    monkeypatch.setattr(upload_module, "_find_existing_document", racing_lookup)

    r = _post_pdf(client)
    assert r.status_code == 202, r.text
    assert r.json()["filename"] == "winner.pdf"
    assert db_session.query(Document).count() == 1
    assert (
        db_session.query(Document).one().content_sha256
        == hashlib.sha256(PDF_BYTES).hexdigest()
    )


def test_read_bounded_returns_all_bytes_under_cap():
    assert _read_bounded(io.BytesIO(b"x" * 10), max_bytes=10) == b"x" * 10


def test_read_bounded_aborts_over_cap_without_full_read():
    class CountingStream:
        def __init__(self, total):
            self.remaining = total
            self.reads = 0

        def read(self, n):
            self.reads += 1
            take = min(n, self.remaining)
            self.remaining -= take
            return b"x" * take

    stream = CountingStream(total=300 * 1024 * 1024)  # pretend 300 MB body
    with pytest.raises(HTTPException) as exc:
        _read_bounded(stream, max_bytes=2 * 1024 * 1024)
    assert exc.value.status_code == 413
    # Aborted after ~3 x 1 MiB reads, not after draining 300 MB.
    assert stream.reads <= 4


def test_oversized_body_413_leaves_no_document_row(client, seeded, db_session, monkeypatch):
    monkeypatch.setattr("routes.upload.MAX_UPLOAD_BYTES", 8)
    files = {"file": ("big.pdf", io.BytesIO(b"0123456789ABCDEF"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 413
    assert db_session.query(Document).count() == 0


def test_storage_write_failure_marks_failed_and_507(client, seeded, db_session, monkeypatch):
    class FailingStore:
        def put(self, key, data):
            raise OSError("disk full")

    monkeypatch.setattr("routes.upload.object_store.get_store", lambda: FailingStore())
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 507
    assert r.json()["detail"]["code"] == "STORAGE_WRITE_FAILED"
    doc = db_session.query(Document).one()
    assert doc.status == "failed"
    assert doc.error is not None


def test_storage_write_failure_507_holds_even_if_mark_failed_commit_raises(
    client, seeded, db_session, monkeypatch
):
    """F-29 hardening: if the mark-failed commit itself raises (e.g. dropped
    connection), the 507 contract must still hold instead of leaking a 500
    and leaving the row stuck 'pending'."""

    class FailingStore:
        def put(self, key, data):
            raise OSError("disk full")

    monkeypatch.setattr("routes.upload.object_store.get_store", lambda: FailingStore())

    real_commit = db_session.commit
    calls = {"n": 0}

    def flaky_commit():
        calls["n"] += 1
        if calls["n"] <= 1:
            # let the rate-limit-usage commit succeed; the pending row is
            # only flushed (not committed) before the write, so the very
            # next commit is the mark-failed one -- that's the one that
            # must fail here.
            return real_commit()
        raise OSError("connection dropped")

    monkeypatch.setattr(db_session, "commit", flaky_commit)

    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 507
    assert r.json()["detail"]["code"] == "STORAGE_WRITE_FAILED"


def test_store_construction_failure_marks_failed_and_507(client, seeded, db_session, monkeypatch):
    """F-29: object_store.get_store() itself (not just store.put) can raise
    -- e.g. bad R2 config -- and must still hit the 507-plus-marked-failed
    contract, not a bare 500."""

    def boom_get_store():
        raise RuntimeError("bad R2 config")

    monkeypatch.setattr("routes.upload.object_store.get_store", boom_get_store)
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 507
    assert r.json()["detail"]["code"] == "STORAGE_WRITE_FAILED"
    doc = db_session.query(Document).one()
    assert doc.status == "failed"
    assert doc.error is not None


def test_upload_rejects_fake_pdf(client, seeded, db_session):
    """F-55: extension says .pdf but bytes are not %PDF -> 415, no row."""
    files = {"file": ("notes.pdf", io.BytesIO(b"MZ\x90\x00 not a pdf"), "application/pdf")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    assert r.status_code == 415
    assert r.json()["detail"]["code"] == "CONTENT_TYPE_MISMATCH"
    assert db_session.query(Document).count() == 0


def test_upload_rejects_fake_pptx(client, seeded):
    files = {
        "file": ("deck.pptx", io.BytesIO(b"%PDF-1.7 wrong container"), "application/octet-stream")
    }
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    assert r.status_code == 415
    assert r.json()["detail"]["code"] == "CONTENT_TYPE_MISMATCH"


def test_upload_accepts_real_pdf_magic(client, seeded):
    files = {"file": ("real.pdf", io.BytesIO(b"%PDF-1.7\n..."), "application/pdf")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    assert r.status_code == 202, r.text


def test_upload_txt_skips_sniff(client, seeded):
    """txt/md have no magic bytes; the sniff must not block them."""
    files = {"file": ("notes.txt", io.BytesIO(b"plain text"), "text/plain")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    assert r.status_code == 202, r.text


def test_upload_writes_through_object_store(client, seeded, db_session, monkeypatch):
    class RecordingStore:
        def __init__(self):
            self.puts = []

        def put(self, key, data):
            self.puts.append((key, data))

    store = RecordingStore()
    monkeypatch.setattr("routes.upload.object_store.get_store", lambda: store)
    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 202
    doc_id = r.json()["document_id"]
    assert store.puts == [(f"{doc_id}_notes.pdf", b"%PDF-fake")]


def test_blob_write_happens_before_pending_row_commit(client, seeded, db_session, monkeypatch):
    """PR-4 Finding 1: the pending row must not be committed (and therefore
    not claimable by the worker's poll loop) until after the blob write
    succeeds. Pins the ordering: put() must run strictly before the commit
    that persists the pending row."""
    order = []

    class RecordingStore:
        def put(self, key, data):
            order.append("put")

    monkeypatch.setattr("routes.upload.object_store.get_store", lambda: RecordingStore())

    real_commit = db_session.commit

    def spy_commit():
        order.append("commit")
        return real_commit()

    monkeypatch.setattr(db_session, "commit", spy_commit)

    files = {"file": ("notes.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post("/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files)
    assert r.status_code == 202, r.text

    # Exactly: rate-limit commit, then the blob write, then the row commit.
    # If the pending row were committed before the write (the race), "put"
    # would land after the second "commit" or the sequence would collapse
    # to ["commit", "commit", "put"].
    assert order == ["commit", "put", "commit"], order


def test_bad_extension_does_not_burn_slot(client, seeded, db_session):
    resp = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files={"file": ("notes.exe", io.BytesIO(b"MZ"), "application/octet-stream")},
    )
    assert resp.status_code == 400
    count = db_session.query(UsageCounter).filter_by(user_id=USER_ID).count()
    assert count == 0


def test_foreign_session_does_not_burn_slot(client, seeded, db_session):
    resp = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": "not_yours"},
        files={"file": ("notes.pdf", io.BytesIO(b"%PDF-1.4 x"), "application/pdf")},
    )
    assert resp.status_code == 404
    count = db_session.query(UsageCounter).filter_by(user_id=USER_ID).count()
    assert count == 0


def test_upload_rejected_when_cost_capped(client, db_session, seeded, monkeypatch):
    monkeypatch.setattr("services.cost_meter.settings.llm_hard_cap_usd", 0.10)

    cost_meter.record_cost(db_session, USER_ID, Decimal("0.2000"))
    db_session.commit()
    files = {"file": ("notes.txt", b"hello", "text/plain")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    assert r.status_code == 429
    assert r.json()["detail"]["code"] == "daily_cost_cap_reached"
    assert "resets_at" in r.json()["detail"]
    count = db_session.query(UsageCounter).filter_by(user_id=USER_ID).count()
    assert count == 0


def test_plaintext_upload_rejected_by_chunk_estimate(client, seeded, monkeypatch):
    monkeypatch.setattr("routes.upload.settings.max_chunks", 10)
    # 10 chunks * 450 tokens * 6.38 chars ~= 28,710 chars; send well past it
    big = b"a" * 60_000
    files = {"file": ("notes.txt", big, "text/plain")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    assert r.status_code == 413
    assert r.json()["detail"]["code"] == "chunk_limit_exceeded"


def test_pdf_upload_not_subject_to_byte_estimate(client, seeded, monkeypatch):
    monkeypatch.setattr("routes.upload.settings.max_chunks", 1)
    files = {"file": ("slides.pdf", b"%PDF-1.4 tiny", "application/pdf")}
    r = client.post(
        "/api/upload",
        data={"user_id": USER_ID, "session_id": SESSION_ID},
        files=files,
    )
    assert r.status_code == 202  # PDFs skip the estimate; worker cap catches them


def _blank_pdf_bytes(pages: int) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_pdf_over_page_limit_413(client, seeded, monkeypatch):
    """F-03: an oversized PDF must be rejected at upload, before a worker ever
    loads and tokenises the whole document."""
    monkeypatch.setattr("routes.upload.settings.max_pages", 2)
    files = {"file": ("big.pdf", io.BytesIO(_blank_pdf_bytes(3)), "application/pdf")}
    r = client.post(
        "/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files
    )
    assert r.status_code == 413, r.text
    detail = r.json()["detail"]
    assert detail["code"] == "page_limit_exceeded"
    assert detail["max_pages"] == 2
    assert detail["page_count"] == 3


def test_pdf_under_page_limit_accepted(client, seeded, monkeypatch):
    monkeypatch.setattr("routes.upload.settings.max_pages", 5)
    files = {"file": ("small.pdf", io.BytesIO(_blank_pdf_bytes(2)), "application/pdf")}
    r = client.post(
        "/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files
    )
    assert r.status_code == 202, r.text


def test_unparseable_pdf_passes_the_page_gate(client, seeded, monkeypatch):
    """The gate is advisory: a file pypdf cannot open is let through so the
    ingestion pipeline reports the extraction failure as it always has."""
    monkeypatch.setattr("routes.upload.settings.max_pages", 1)
    files = {"file": ("broken.pdf", io.BytesIO(b"%PDF-fake"), "application/pdf")}
    r = client.post(
        "/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files
    )
    assert r.status_code == 202, r.text


def test_pptx_over_slide_limit_413(client, seeded, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr("routes.upload.settings.max_pages", 1)
    prs = Presentation()
    for _ in range(2):
        prs.slides.add_slide(prs.slide_layouts[6])
    buf = io.BytesIO()
    prs.save(buf)
    files = {
        "file": (
            "deck.pptx",
            io.BytesIO(buf.getvalue()),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
    }
    r = client.post(
        "/api/upload", data={"user_id": USER_ID, "session_id": SESSION_ID}, files=files
    )
    assert r.status_code == 413, r.text
    detail = r.json()["detail"]
    assert detail["code"] == "page_limit_exceeded"
    assert detail["page_count"] == 2
